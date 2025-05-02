import os
import re
from docx import Document
from openpyxl import Workbook

# 读取 docx 文件
def read_docx(file_path):
    doc = Document(file_path)
    full_text = []
    for para in doc.paragraphs:
        full_text.append(para.text)
    return "\n".join(full_text)

# 按照句号和问号拆分文本
def split_text(text):
    # 使用正则表达式按句号（。）、问号（?）进行拆分，并保留标点符号
    sentences = re.split(r'([。?])', text)  # 保留标点符号
    # 合并标点符号到前一个句子
    sentences = [sentences[i] + sentences[i+1] if i+1 < len(sentences) else sentences[i] 
                 for i in range(0, len(sentences), 2)]
    return [s.strip() for s in sentences if s.strip()]

# 保存拆分后的文本到 Excel 文件
def save_to_excel(text_list, output_file):
    wb = Workbook()
    ws = wb.active
    ws.title = "拆分内容"
    
    # 将每条文本写入Excel的一行
    for i, sentence in enumerate(text_list, start=1):
        ws[f'A{i}'] = sentence
    
    # 保存文件
    wb.save(output_file)

# 主程序
def main():
    # 定义文件路径
    docx_file = r"C:\Users\17439\Desktop\毕业论文\朝春小学教研\【教研】-编码数据\2024教研下\王老师+数学.docx"
    folder = os.path.dirname(docx_file)  # 获取文件所在目录
    output_file = os.path.join(folder, "拆分结果.xlsx")  # 新文件路径

    # 读取 docx 文件
    text = read_docx(docx_file)

    # 拆分文本
    sentences = split_text(text)

    # 保存拆分内容到 Excel 文件
    save_to_excel(sentences, output_file)
    
    print(f"拆分内容已保存到: {output_file}")

if __name__ == "__main__":
    main()
